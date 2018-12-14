from docx import Document
from openpyxl import Workbook
import glob

from openpyxl.cell import Cell
from openpyxl.styles import Font

wb = Workbook()

ws = wb.active

# add title
ws.append([
    '问题编号（填写文件名）',
    '渠道',
    '问题详细描述',
    '问题类型',
    '问题解决方案',
    '责任人',
    '预计解决时间',
    '解决状态',
    '实际解决日期'
])


def styled_cells(data):
    for idx, val in enumerate(data):
        if idx == 0:
            val = Cell(ws, value=val)
            val.font = Font(underline='single', color='0563C1')
        yield val


for doc in glob.iglob("*.docx"):

    document = Document(doc)

    # 唯一表格
    table = document.tables[0]

    # 第五行，渠道
    client = table.rows[4]

    # 第六行，问题详细描述
    desc = table.rows[5]

    ws.append(styled_cells([
        '=HYPERLINK("{}", "{}")'.format(doc, doc[0:-5]),
        client.cells[1].text.strip(),
        desc.cells[1].text.strip()
    ]))


# Save the file
wb.save("sample.xlsx")
